"""Generate synthetic examples and reproducible reports, without downloads."""
from __future__ import annotations

import argparse
import io
from pathlib import Path

from .cli import main as check

EXAMPLE = """# Synthetic imaging manuscript — formatting demonstration only

## Abstract
Background: This is a synthetic example.
Methods: We generated n = 12 fictional observations.
Results: A t-test gave P = 0.000.
Conclusion: These are deliberately flawed reporting examples.

## Keywords
synthetic example, manuscript checklist

## Introduction
This demonstration contains no patient data and supports no scientific claim.

## Methods
All observations and reference entries are fictional. A t-test is mentioned for rule demonstration only.

## Results
We deliberately omit uncertainty estimates and any scientific interpretation.

## Discussion
Authors must verify reporting requirements and ethics applicability themselves.

## References
[1] Example A. Fictional reference for demonstration. 2026.
[2] Example A. Fictional reference for demonstration. 2026.
[4] Example B. Another fictional reference. 2026.
[4] Example C. Repeated number label example. 2026.

## Funding
No funding was received for this synthetic demonstration.

## Conflicts of interest
None declared for this synthetic demonstration.

## Data availability
No patient or study data exist for this demonstration.
"""


def main(argv=None):
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", default="demo-output", help="Destination for NEW example files (default: demo-output)")
    args = parser.parse_args(argv)
    output = Path(args.out_dir)
    paths = [output / name for name in ("synthetic_manuscript.md", "synthetic_manuscript.docx", "report.md")]
    if any(path.exists() for path in paths):
        parser.error("Example files already exist; choose a new --out-dir to preserve them.")
    output.mkdir(parents=True, exist_ok=True)
    paths[0].write_text(EXAMPLE, encoding="utf-8")
    document = Document()
    for line in EXAMPLE.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], 0)
        elif line.startswith("## "):
            document.add_heading(line[3:], 1)
        else:
            document.add_paragraph(line)
    image_blob = io.BytesIO()
    Image.new("RGB", (120, 120), "#4169e1").save(image_blob, format="PNG")
    image_blob.seek(0)
    document.add_picture(image_blob, width=Inches(2), height=Inches(2))
    document.save(paths[1])
    print("Synthetic demonstration: no patient data or scientific claims.\n")
    return check([str(paths[1]), "--language", "en", "--md", str(paths[2])])


if __name__ == "__main__":
    raise SystemExit(main())
