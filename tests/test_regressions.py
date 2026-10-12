"""Regression tests for reported defects and article-profile boundaries."""
import contextlib
import io
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.shared import Inches
from PIL import Image

from manulint.checks import (CheckConfig, Finding, _section_lines, check_ethics,
                             check_figures, check_language, check_meta,
                             check_references, check_statistics, check_structure,
                             run_all, score)
from manulint.cli import main
from manulint.readers import Manuscript, _extract_docx_figures, read_manuscript
from manulint.report import render_markdown, render_terminal


def rules(findings):
    return {finding.rule for finding in findings}


class RegressionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def text_file(self, text, name="paper.md"):
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path

    def invoke(self, argv):
        stdout, stderr = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            result = main(argv)
        return result, stdout.getvalue(), stderr.getvalue()

    def test_empty_text_is_input_error_without_score(self):
        for text in ("", "\n\t   \n"):
            result, output, error = self.invoke([str(self.text_file(text))])
            self.assertEqual(result, 2)
            self.assertNotIn("score", output)
            self.assertIn("no extractable text", error)
        with self.assertRaises(ValueError):
            run_all(Manuscript(" "))

    def test_empty_docx_is_input_error(self):
        path = self.root / "blank.docx"
        Document().save(path)
        result, output, error = self.invoke([str(path)])
        self.assertEqual(result, 2)
        self.assertNotIn("score", output)
        self.assertIn("no extractable text", error)

    def test_corrupt_docx_reports_clear_error_without_traceback(self):
        path = self.text_file("not a zip", "bad.docx")
        result, output, error = self.invoke([str(path)])
        self.assertEqual(result, 2)
        self.assertIn("Cannot read DOCX", error)
        self.assertNotIn("Traceback", error)

    def test_abstract_word_count_with_blank_lines(self):
        text = "Title\n\n## Abstract\n\none two three\n\n## Introduction\nfour five"
        finding = next(f for f in check_meta(Manuscript(text)) if f.rule == "abstract-words")
        self.assertIn("3 whitespace-delimited words", finding.message)

    def test_structured_abstract_labels_not_body_sections_or_words(self):
        text = "Title\nAbstract\nBackground: one two\nMethods\nthree four\nResults: five\nConclusion: six\nKeywords: example\nIntroduction\nBody.\nReferences\n[1] Example."
        sections = _section_lines(text)
        self.assertNotIn("methods", sections)
        self.assertNotIn("results", sections)
        finding = next(f for f in check_meta(Manuscript(text)) if f.rule == "abstract-words")
        self.assertIn("6 whitespace-delimited words", finding.message)
        self.assertIn("missing-methods", rules(check_structure(Manuscript(text))))

    def test_inline_abstract_and_chinese_character_count(self):
        manuscript = Manuscript("题目\n摘要：目的：研究儿童\n方法：影像检查\n关键词：磁共振\n引言\n正文。")
        # Inline structured label is content of an inline Abstract, not a heading.
        finding = next(f for f in check_meta(manuscript, CheckConfig(language="zh")) if f.rule == "abstract-characters")
        self.assertIn("8 Chinese characters", finding.message)

    def test_body_mentions_are_not_headings(self):
        text = "Title\nOur results were promising.\nThe methods are described below.\nThe abstract discussion is brief."
        self.assertEqual(_section_lines(text), {})
        self.assertIn("missing-results", rules(check_structure(Manuscript(text))))
        self.assertEqual(_section_lines("## 2.1 Materials and Methods\n**References**"), {"methods": 0, "references": 1})

    def test_reference_number_free_duplicate_detection(self):
        text = "References\n[1] Smith J. Same reference. 2024.\n[2] Smith J. Same reference. 2024."
        result = rules(check_references(Manuscript(text)))
        self.assertIn("duplicates", result)
        self.assertNotIn("number-duplicate", result)

    def test_reference_gap_duplicate_label_and_order_are_distinct(self):
        text = "References\n1. A\n3. B\n3. C\n2. D\n5. E"
        result = rules(check_references(Manuscript(text)))
        self.assertIn("number-gap", result)
        self.assertIn("number-duplicate", result)
        self.assertIn("number-order", result)
        self.assertNotIn("ok", result)

    def test_reference_check_stops_at_following_section(self):
        text = "References\n[1] A\n[2] B\nAcknowledgments\n[4] C\n[4] C"
        self.assertEqual(rules(check_references(Manuscript(text))), {"ok"})

    def test_author_year_duplicates_are_checked(self):
        text = "References\nSmith, J. (2024). A.\nSmith, J. (2024). A."
        self.assertEqual(rules(check_references(Manuscript(text))), {"unparsed", "duplicates"})

    def test_article_type_profiles_do_not_force_imrad_on_review(self):
        text = "Title\nAbstract\nSummary.\nIntroduction\nText.\nReferences\n[1] A"
        result = rules(check_structure(Manuscript(text), CheckConfig(article_type="review")))
        self.assertNotIn("missing-methods", result)
        self.assertNotIn("missing-results", result)
        case = rules(check_structure(Manuscript(text), CheckConfig(article_type="case-report")))
        self.assertIn("missing-case_presentation", case)
        self.assertNotIn("missing-methods", case)
        self.assertEqual(rules(check_statistics(Manuscript(text), CheckConfig(article_type="case-report"))), {"profile-skip"})

    def test_p_style_case_is_configured_not_universal(self):
        manuscript = Manuscript("n = 12; t-test; P = 0.03; p = 0.04; 95% CI")
        self.assertNotIn("p-case", rules(check_statistics(manuscript)))
        self.assertIn("p-case", rules(check_statistics(manuscript, CheckConfig(p_style="lower"))))
        self.assertIn("p-case", rules(check_statistics(manuscript, CheckConfig(p_style="upper"))))
        self.assertIn("p-zero", rules(check_statistics(Manuscript("p = 0.000"))))

    def test_chinese_punctuation_allowed_in_chinese_and_auto_profiles(self):
        manuscript = Manuscript("这是中文稿件，包含中文标点。")
        self.assertNotIn("cjk-punct", rules(check_language(manuscript)))
        self.assertNotIn("cjk-punct", rules(check_language(manuscript, CheckConfig(language="zh"))))
        self.assertIn("cjk-punct", rules(check_language(manuscript, CheckConfig(language="en"))))

    def test_ethics_keywords_do_not_claim_adequacy(self):
        findings = check_ethics(Manuscript("IRB approval, informed consent"))
        self.assertEqual(rules(findings), {"not-validated"})
        self.assertIn("cannot establish ethical adequacy", findings[0].message)

    def test_score_and_reports_never_claim_publication_readiness(self):
        value, label = score([])
        report = render_terminal([], value, label)
        self.assertEqual(value, 100)
        self.assertIn("unvalidated", report)
        self.assertIn("does not measure publication readiness", report)
        self.assertNotIn("可投稿", report)

    def make_figure_docx(self):
        image_a, image_b = self.root / "a.png", self.root / "b.png"
        Image.new("RGB", (600, 300), "white").save(image_a)
        Image.new("RGB", (120, 120), "blue").save(image_b)
        doc = Document()
        doc.add_paragraph("Synthetic figure test")
        doc.add_picture(str(image_a), width=Inches(1), height=Inches(0.5))
        doc.add_picture(str(image_b), width=Inches(1), height=Inches(1))
        doc.add_picture(str(image_a), width=Inches(1), height=Inches(4))
        path = self.root / "figures.docx"
        doc.save(path)
        return path

    def test_reused_image_occurrences_match_rids_and_both_dimensions(self):
        manuscript = read_manuscript(self.make_figure_docx())
        self.assertEqual(len(manuscript.figures), 3)
        self.assertEqual([figure.image.size for figure in manuscript.figures], [(600, 300), (120, 120), (600, 300)])
        findings = check_figures(manuscript)
        self.assertEqual([f.severity for f in findings], ["info", "fail", "fail"])
        self.assertIn("75 dpi", findings[2].message)
        self.assertEqual(manuscript.figures[2].height_in, 4)

    def test_image_matching_independent_of_zip_member_order(self):
        path = self.make_figure_docx()
        reordered = self.root / "reordered.docx"
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(reordered, "w") as target:
            for name in reversed(source.namelist()):
                target.writestr(name, source.read(name))
        self.assertEqual([f.image.size for f in read_manuscript(reordered).figures], [(600, 300), (120, 120), (600, 300)])

    def test_vector_images_are_explicitly_skipped(self):
        doc = Document(self.make_figure_docx())
        related_id = doc.element.body.xpath(".//a:blip")[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        doc.part.related_parts[related_id] = SimpleNamespace(partname="/word/media/vector.svg", blob=b"<svg/>")
        figures = _extract_docx_figures(doc)
        self.assertIsNone(figures[0].image)
        self.assertIn("Vector format", figures[0].skip_reason)
        self.assertIn("skipped", rules(check_figures(Manuscript("text", figures))))

    def test_docx_paragraphs_and_tables_keep_body_order(self):
        doc = Document()
        doc.add_paragraph("Before")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "In table"
        doc.add_paragraph("After")
        path = self.root / "table.docx"
        doc.save(path)
        self.assertEqual(read_manuscript(path).text, "Before\nIn table\nAfter")

    def test_markdown_escapes_all_table_fields_and_source(self):
        finding = Finding("warn", "a|b", "rule|x", "<script>|message\n**strong**", "`evidence`|<b>")
        report = render_markdown([finding], 99, "label|", "paper`\n# forged.md")
        self.assertNotIn("<script>", report)
        self.assertNotIn("\n# forged", report)
        self.assertIn("a\\|b", report)
        self.assertIn("&lt;script&gt;\\|message", report)
        self.assertNotIn("**strong**", report)

    def test_cli_report_never_overwrites_source(self):
        path = self.text_file("Keep manuscript")
        result, output, error = self.invoke([str(path), "--md", str(path)])
        self.assertEqual(result, 2)
        self.assertEqual(path.read_text(), "Keep manuscript")
        self.assertIn("must differ", error)

    def test_hardlink_report_path_cannot_overwrite_manuscript(self):
        source = self.text_file("Keep source manuscript")
        output = self.root / "hardlink-report.md"
        os.link(source, output)
        result, stdout, stderr = self.invoke([str(source), "--md", str(output)])
        self.assertEqual(result, 2)
        self.assertIn("must differ", stderr)
        self.assertEqual(source.read_text(), "Keep source manuscript")

    def test_symlink_report_path_cannot_overwrite_manuscript(self):
        source = self.text_file("Keep source manuscript")
        output = self.root / "symlink-report.md"
        try:
            output.symlink_to(source)
        except (OSError, NotImplementedError):
            self.skipTest("Symlink creation is unavailable for this account")
        result, stdout, stderr = self.invoke([str(source), "--md", str(output)])
        self.assertEqual(result, 2)
        self.assertIn("must differ", stderr)
        self.assertEqual(source.read_text(), "Keep source manuscript")

    def test_existing_report_is_replaced_and_source_preserved(self):
        source = self.text_file("Title\np = 0.000")
        output = self.text_file("old report", "report.md")
        result, stdout, stderr = self.invoke([str(source), "--md", str(output)])
        self.assertEqual(result, 0)
        self.assertIn("manulint rule-check report", output.read_text())
        self.assertEqual(source.read_text(), "Title\np = 0.000")
        self.assertEqual(list(self.root.glob(".manulint-*.tmp")), [])

    def test_invalid_cta_returns_input_error(self):
        source = self.text_file("Title")
        result, stdout, stderr = self.invoke([str(source), "--cta", "http://[broken"])
        self.assertEqual(result, 2)
        self.assertIn("Input error", stderr)

    def test_cli_exit_zero_means_completed_not_passed(self):
        result, output, error = self.invoke([str(self.text_file("Title\np = 0.000")), "--language", "en"])
        self.assertEqual(result, 0)
        self.assertIn("[FLAG]", output)
        self.assertIn("unvalidated", output)
        self.assertEqual(error, "")


if __name__ == "__main__":
    unittest.main()
